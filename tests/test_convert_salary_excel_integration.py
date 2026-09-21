"""Exercise the documented Excel -> converter CLI -> salary JSON path."""
import json
import subprocess
import sys
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

try:
    import openpyxl
except ImportError:
    openpyxl = None

from salary_lookup import collect_validation_issues, format_entry

TOOL = Path(__file__).resolve().parent.parent / "tools" / "convert_salary_excel.py"


@unittest.skipUnless(openpyxl, "requires optional openpyxl (installed in CI)")
class ConverterCLI(unittest.TestCase):
    def setUp(self):
        directory = TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.workbook = self.root / "salary input.xlsx"
        self.output = self.root / "salary output.json"

    def run_converter(self, *args):
        return subprocess.run(
            [sys.executable, str(TOOL), str(self.workbook),
             "--output", str(self.output), *args],
            cwd=self.root, capture_output=True, text=True,
        )

    def test_real_workbook_round_trips_multiple_sheets_and_metadata(self):
        book = openpyxl.Workbook()
        try:
            first = book.active
            first.title = "All employees"
            first.append(["Company", "City", "Count", "Index"])
            first.append(["Acme Corp", "Copenhagen", 500, 108.5])
            first.append(["Café Étoile", "Paris", 12, "102,5"])
            second = book.create_sheet("Engineering")
            second.append(["Company", "City", "Engineering Count", "Engineering Index"])
            second.append(["Other Corp", "Berlin", 8, 112.5])
            book.save(self.workbook)
        finally:
            book.close()

        result = self.run_converter(
            "--source", "Integration fixture", "--baseline", "100",
            "--baseline-desc", "Index 100 = median salary",
        )
        self.assertEqual(result.returncode, 0, result.stderr)
        data = json.loads(self.output.read_text(encoding="utf-8"))
        self.assertEqual(data["metadata"], {
            "source": "Integration fixture", "index_baseline": 100,
            "index_label": "Index", "baseline_description": "Index 100 = median salary",
        })
        self.assertEqual(data["companies"], [
            {"company": "Acme Corp", "city": "Copenhagen", "categories": {
                "all_employees": {"count": 500, "index": 108.5}}},
            {"company": "Café Étoile", "city": "Paris", "categories": {
                "all_employees": {"count": 12, "index": 102.5}}},
            {"company": "Other Corp", "city": "Berlin", "categories": {
                "engineering": {"count": 8, "index": 112.5}}},
        ])
        self.assertEqual(collect_validation_issues(data), ([], []))
        rendered = format_entry(data["companies"][0], data["metadata"])
        self.assertIn("All Employees", rendered)
        self.assertIn("+8.5%", rendered)
        self.assertNotIn("N/A*", rendered)

    def test_workbook_without_salary_headers_fails_without_writing_output(self):
        book = openpyxl.Workbook()
        try:
            book.active.append(["Notes", "Comments"])
            book.active.append(["No salary data", "Just a worksheet"])
            book.save(self.workbook)
        finally:
            book.close()

        result = self.run_converter()
        self.assertEqual(result.returncode, 1, result.stderr)
        self.assertIn("No data could be parsed", result.stderr)
        self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
